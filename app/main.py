import hashlib
import json
from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

from flask import (
    Blueprint,
    current_app,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from flask_login import current_user, login_required
from werkzeug.utils import secure_filename

from app import db
from app.ai.backup import build_backup
from app.ai.assessment import build_assessment
from app.ai.catalog import get_algorithm_catalog
from app.ai.engine import analyze_inspection
from app.ai.problem_solver import solve_visual_problems
from app.ai.training import build_training_samples, train_visual_model
from app.ai.types import InspectionInputs
from app.ai.vision import (
    aggregate_inspection_fields,
    aggregate_photo_models,
    analyze_available_models,
    assess_image,
)
from app.models import (
    DiseaseObservation,
    FieldPrediction,
    FloweringObservation,
    FruitObservation,
    GrowthMeasurement,
    AgronomicAssessment,
    ModelMetadata,
    ObservationSession,
    PestObservation,
    Photo,
    SoilObservation,
    Tree,
    UserCorrection,
    WeedObservation,
)

bp = Blueprint("main", __name__)


@bp.get("/about")
@login_required
def about():
    algorithm_groups = [
        {
            "title": "Analisis foto dan computer vision",
            "items": [
                (
                    "Quality Gate",
                    "Memeriksa format, resolusi, pencahayaan, dan kelayakan foto.",
                ),
                ("YOLO11 Nano", "Mendeteksi objek umum pada foto."),
                (
                    "MobileNetV3 Classification",
                    "Membaca kelas visual umum dari gambar.",
                ),
                ("MobileNetV3 Embedding", "Mengekstrak representasi fitur visual."),
                ("DeepLabV3 MobileNetV3", "Membuat ringkasan segmentasi area visual."),
                (
                    "Feature Engineering",
                    "Menghitung brightness, warna, dominasi hijau, dan ketajaman.",
                ),
            ],
        },
        {
            "title": "Analisis pertanian",
            "items": [
                (
                    "Flower/Fruit Screening",
                    "Menyaring keberadaan bunga, buah, dan perkiraan jumlah objek.",
                ),
                (
                    "Pest/Disease/Weed Screening",
                    "Menggabungkan deteksi, fitur visual, segmentasi, dan bukti anomali.",
                ),
                (
                    "Virtual Soil Sensor",
                    "Memperkirakan kelas kondisi zona akar dari foto dan riwayat.",
                ),
                (
                    "Tree Structure Screening",
                    "Menyaring struktur tajuk, vigor, dan visibilitas pohon.",
                ),
                (
                    "Stem/Branch Screening",
                    "Menyaring visibilitas batang dan kerusakan visual.",
                ),
                (
                    "Shoot Growth Screening",
                    "Menyaring tunas dan menyiapkan analisis pertumbuhan historis.",
                ),
                (
                    "Logistic Regression",
                    "Klasifikasi kondisi sederhana ketika data berlabel sudah cukup.",
                ),
                (
                    "Decision Tree",
                    "Membuat keputusan klasifikasi yang mudah dijelaskan.",
                ),
                (
                    "Random Forest",
                    "Klasifikasi dan estimasi risiko berbasis banyak pohon keputusan.",
                ),
                (
                    "Gradient Boosting / XGBoost-compatible",
                    "Meningkatkan prediksi klasifikasi dan risiko secara bertahap.",
                ),
                (
                    "SVM / One-Class SVM",
                    "Klasifikasi dan deteksi kondisi yang menyimpang dari pola normal.",
                ),
                (
                    "K-Nearest Neighbors",
                    "Membandingkan kondisi pohon dengan contoh historis terdekat.",
                ),
                (
                    "Naive Bayes",
                    "Klasifikasi probabilistik untuk data gejala dan kondisi lapangan.",
                ),
                (
                    "Linear Regression",
                    "Estimasi pertumbuhan dari fitur pengukuran yang tersedia.",
                ),
                (
                    "Random Forest Regression",
                    "Regresi non-linear untuk pertumbuhan, risiko, dan hasil panen.",
                ),
                (
                    "Isolation Forest",
                    "Mendeteksi observasi pertumbuhan atau kesehatan yang tidak biasa.",
                ),
                (
                    "Statistical / Historical Deviation",
                    "Membandingkan nilai baru dengan baseline individual setiap pohon.",
                ),
                (
                    "K-Means dan DBSCAN",
                    "Mengelompokkan pola kesehatan, pertumbuhan, dan kondisi kebun.",
                ),
                (
                    "PCA dan Embedding Reduction",
                    "Meringkas dimensi fitur visual dan historis untuk analisis pola.",
                ),
                (
                    "Moving Average, EWMA, dan Rolling Statistics",
                    "Menghaluskan noise dan membaca perubahan kondisi dari waktu ke waktu.",
                ),
                (
                    "Trend Slope dan ARIMA-style Forecasting",
                    "Menganalisis arah serta prakiraan pertumbuhan jika histori cukup.",
                ),
                (
                    "LSTM / GRU",
                    "Fondasi prediksi time-series jangka panjang setelah dataset memadai.",
                ),
                (
                    "RT-DETR, YOLO dan Object Detection",
                    "Mendeteksi pohon, daun, bunga, buah, batang, dan objek lapangan.",
                ),
                (
                    "SAM / MobileSAM dan U-Net",
                    "Segmentasi tajuk, daun, buah, tanah, dan area kerusakan.",
                ),
                (
                    "Depth Estimation dan Reference Measurement",
                    "Estimasi ukuran dengan depth, marker, ruler, atau objek referensi.",
                ),
                (
                    "Weighted Multi-Model Ensemble",
                    "Menggabungkan model lokal, rule, histori, dan bukti visual.",
                ),
                (
                    "SHAP, LIME, Grad-CAM, dan Heatmap",
                    "Menjelaskan fitur, area gambar, dan alasan di balik keluaran model.",
                ),
                (
                    "Risk Prediction dan Growth Prediction",
                    "Mengubah pola lapangan menjadi skor risiko dan arah pertumbuhan.",
                ),
                (
                    "Calibration, Agreement, dan Out-of-Distribution",
                    "Menurunkan confidence saat model tidak sepakat atau data di luar pola.",
                ),
            ],
        },
        {
            "title": "Machine learning dan analisis historis",
            "items": [
                (
                    "Rule-Based Expert System",
                    "Menghitung skor kesehatan dan risiko dari aturan yang transparan.",
                ),
                (
                    "Random Forest / Gradient Boosting",
                    "Kandidat klasifikasi saat label cukup; tidak aktif tanpa data minimum.",
                ),
                (
                    "Anomaly Detection",
                    "Membandingkan kondisi terbaru dengan baseline pohon.",
                ),
                (
                    "Time-Series Trend",
                    "Membaca tren pertumbuhan dan perubahan historis.",
                ),
                (
                    "Weighted Ensemble",
                    "Menggabungkan keluaran beberapa model dengan confidence.",
                ),
                (
                    "Training dari koreksi pengguna",
                    "Menggunakan koreksi tervalidasi sebagai label, bukan prediksi mentah.",
                ),
            ],
        },
        {
            "title": "Explainable AI dan keselamatan data",
            "items": [
                (
                    "XAI Evidence",
                    "Menampilkan alasan, sumber, confidence, dan provenance.",
                ),
                (
                    "Confidence Scoring",
                    "Menurunkan keyakinan saat model tidak sepakat atau data kurang.",
                ),
                (
                    "Risk and Recommendation",
                    "Menghasilkan skor risiko dan rekomendasi berdasarkan bukti.",
                ),
                (
                    "Human-in-the-Loop",
                    "Meminta konfirmasi untuk nilai yang belum dapat ditentukan.",
                ),
                (
                    "Insufficient-data guard",
                    "Tidak membuat angka, diagnosis, atau prediksi tanpa bukti cukup.",
                ),
            ],
        },
    ]
    return render_template(
        "about.html",
        algorithm_groups=algorithm_groups,
        algorithm_catalog=get_algorithm_catalog(),
    )


def build_ml_form_fields(analysis_payload: dict | None) -> list[dict]:
    """Turn pipeline output into safe, user-facing form fields."""
    if not isinstance(analysis_payload, dict):
        return []
    fields = analysis_payload.get("auto_fields", {})
    labels = {
        "leaf_color_observed": "Warna visual daun",
        "visual_brightness_mean": "Brightness rata-rata",
        "visual_sharpness_proxy": "Indikator ketajaman",
        "detected_objects": "Objek terdeteksi",
        "generic_image_class": "Klasifikasi visual umum",
        "segmentation_summary": "Ringkasan segmentasi",
        "root_zone_moisture_class": "Kelas kelembapan zona akar",
        "pest_or_disease": "Hama atau penyakit",
    }
    form_fields = {
        "surface_dark",
        "standing_water",
        "leaf_wilt",
        "soil_moisture_visual",
        "root_zone_moisture_class",
        "height_cm",
        "stem_diameter_cm",
        "canopy_width_cm",
        "flower_present",
        "flower_stage",
        "fruit_present",
        "fruit_stage",
        "fruit_count_estimate",
        "fruit_damage_percent",
        "pest_present",
        "pest_type",
        "disease_present",
        "disease_type",
        "weed_level",
        "weed_coverage_percent",
        "notes",
    }
    result = []
    for key, field in fields.items():
        if key in form_fields:
            continue
        value = field.get("value") if isinstance(field, dict) else field
        if key == "detected_objects" and value == []:
            value = "Tidak ada objek terdeteksi"
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        result.append(
            {
                "key": key,
                "label": labels.get(key, key.replace("_", " ").title()),
                "value": value if value not in (None, "") else "Belum dapat ditentukan",
                "status": (
                    field.get("status", "NEEDS_CONFIRMATION")
                    if isinstance(field, dict)
                    else "NEEDS_CONFIRMATION"
                ),
                "confidence": (
                    float(field.get("confidence", 0.0))
                    if isinstance(field, dict)
                    else 0.0
                ),
                "source": (
                    field.get("source", "UNKNOWN")
                    if isinstance(field, dict)
                    else "UNKNOWN"
                ),
            }
        )
    return result


def _analysis_summary(results):
    summary = {}
    for key, result in results.items():
        quality = result.get("quality")
        summary[key] = {
            "status": result.get("status"),
            "model": result.get("model"),
            "message": result.get("message"),
        }
        if quality:
            summary[key]["quality_status"] = quality.status
            summary[key]["quality_reasons"] = list(quality.reasons)
        if key == "object_detection" and result.get("status") == "READY":
            summary[key]["detections"] = result.get("detections", [])
        if key == "classification" and result.get("status") == "READY":
            summary[key]["class_id"] = result.get("class_id")
            summary[key]["class_label"] = result.get("class_label")
            summary[key]["confidence"] = result.get("confidence")
        if key == "embedding" and result.get("status") == "READY":
            summary[key]["dimensions"] = result.get("dimensions")
        if key == "segmentation" and result.get("status") == "READY":
            summary[key]["class_count"] = result.get("class_count")
            summary[key]["pixel_count"] = result.get("pixel_count")
            summary[key]["class_distribution"] = result.get("class_distribution")
        if key == "visual_features" and result.get("status") == "READY":
            for field in (
                "brightness_mean",
                "green_mean",
                "green_dominance",
                "sharpness_proxy",
                "color_class",
                "confidence",
            ):
                summary[key][field] = result.get(field)
    return summary


def _analyze_uploaded_photo(file_storage):
    original = secure_filename(file_storage.filename or "")
    extension = original.rsplit(".", 1)[-1].lower() if "." in original else ""
    if extension not in current_app.config["ALLOWED_IMAGE_EXTENSIONS"]:
        return {"status": "INVALID", "message": "Format foto tidak didukung."}
    temporary_folder = Path(current_app.config["UPLOAD_FOLDER"]) / "_analysis"
    temporary_folder.mkdir(parents=True, exist_ok=True)
    temporary_path = (
        temporary_folder
        / f"preview-{datetime.now(timezone.utc).timestamp()}.{extension}"
    )
    try:
        file_storage.save(temporary_path)
        quality = assess_image(temporary_path)
        if quality.status != "GOOD":
            return {
                "status": quality.status,
                "quality_status": quality.status,
                "quality_reasons": list(quality.reasons),
                "models": {},
                "message": "Inference dihentikan sampai foto diperbaiki.",
            }
        models = analyze_available_models(
            temporary_path,
            current_app.config.get(
                "MODEL_FOLDER",
                Path(current_app.root_path).parent / "instance" / "models",
            ),
        )
        return {
            "status": "READY",
            "quality_status": quality.status,
            "quality_reasons": list(quality.reasons),
            "models": _analysis_summary(models),
            "message": (
                "Ekstraksi selesai. Data agronomi yang tidak didukung model "
                "ditandai perlu konfirmasi."
            ),
        }
    finally:
        temporary_path.unlink(missing_ok=True)


@bp.post("/api/trees/<int:tree_id>/inspection/analyze")
@login_required
def analyze_inspection_photo(tree_id):
    db.get_or_404(Tree, tree_id)
    files = [file for file in request.files.getlist("photos") if file.filename]
    if not files:
        return (
            jsonify({"status": "INVALID", "message": "Pilih minimal satu foto."}),
            400,
        )
    analyses = [_analyze_uploaded_photo(file) for file in files]
    combined_models = aggregate_photo_models(analyses)
    aggregated = aggregate_inspection_fields(combined_models)
    problem_analysis = solve_visual_problems(combined_models)
    algorithms = [
        {
            "name": "Quality Gate",
            "status": "READY",
            "details": "Resolusi, pencahayaan, dan validitas file",
        },
        {
            "name": "Feature Engineering",
            "status": "READY",
            "details": "Warna, brightness, dominasi hijau, ketajaman",
        },
        {
            "name": "Weighted Ensemble",
            "status": (
                "READY" if aggregated["ready_model_count"] > 1 else "INSUFFICIENT_DATA"
            ),
            "details": "Penggabungan keluaran model lokal",
        },
        {
            "name": "Anomaly Detection",
            "status": "INSUFFICIENT_DATA",
            "details": "Baseline historis per pohon belum cukup",
        },
        {
            "name": "Agricultural Classifier",
            "status": (
                "READY"
                if aggregated["fields"]["pest_or_disease"]["source"]
                == "VISUAL_ANOMALY_SCREENING"
                else "NEEDS_MODEL"
            ),
            "details": (
                "Penyaringan anomali visual aktif; bukan diagnosis hama/penyakit."
                if aggregated["fields"]["pest_or_disease"]["source"]
                == "VISUAL_ANOMALY_SCREENING"
                else "Dataset rambutan berlabel belum tersedia"
            ),
        },
    ]
    confidence_values = [
        field.get("confidence", 0.0)
        for field in aggregated["fields"].values()
        if field.get("status") not in {"NEEDS_CONFIRMATION", "UNKNOWN"}
    ]
    overall_confidence = (
        sum(confidence_values) / len(confidence_values) if confidence_values else 0.0
    )
    uncertain_count = sum(
        field.get("status") == "NEEDS_CONFIRMATION"
        for field in aggregated["fields"].values()
    )
    analyst_summary = {
        "overall_confidence": round(overall_confidence, 4),
        "uncertain_fields": uncertain_count,
        "priority": ("PERLU KONFIRMASI" if uncertain_count else "PEMANTAUAN NORMAL"),
        "recommendation": (
            "Periksa field yang belum memiliki model pertanian khusus."
            if uncertain_count
            else "Data visual siap ditinjau dan disimpan."
        ),
    }
    agronomic_assessment = build_assessment(
        aggregated["fields"], photo_count=len(files)
    )
    return jsonify(
        {
            "status": "READY",
            "photo_count": len(analyses),
            "analyses": analyses,
            "auto_fields": aggregated["fields"],
            "problem_analysis": problem_analysis,
            "algorithms": algorithms,
            "analyst_summary": analyst_summary,
            "agronomic_assessment": agronomic_assessment,
            "pipeline": {
                "ready_model_count": aggregated["ready_model_count"],
                "model_count": aggregated["model_count"],
                "message": "Semua model terpasang dijalankan; field tanpa model khusus perlu konfirmasi.",
            },
        }
    )


@bp.get("/")
def index():
    if current_user.is_authenticated:
        return redirect(url_for("main.dashboard"))
    return redirect(url_for("auth.login"))


@bp.get("/dashboard")
@login_required
def dashboard():
    trees = db.session.scalars(
        db.select(Tree).where(Tree.active.is_(True)).order_by(Tree.code)
    ).all()
    observations = (
        db.session.scalar(db.select(db.func.count(ObservationSession.id))) or 0
    )
    now = datetime.now(timezone.utc)
    cutoff = now - timedelta(days=30)
    tree_rows = []
    confidence_values = []
    daily_counts = Counter()
    status_counts = Counter()
    priority_rows = []
    for tree in trees:
        ordered = sorted(
            tree.observations,
            key=lambda item: item.observation_datetime,
            reverse=True,
        )
        latest = ordered[0] if ordered else None
        predictions = latest.predictions if latest else []
        latest_confidence = (
            sum(item.confidence for item in predictions) / len(predictions)
            if predictions
            else 0.0
        )
        uncertain = sum(item.confidence < 0.45 for item in predictions)
        status = tree.status or "BELUM DIANALISIS"
        status_counts[status] += 1
        if latest_confidence:
            confidence_values.append(latest_confidence)
        if latest:
            observation_time = latest.observation_datetime
            if observation_time.tzinfo is None:
                observation_time = observation_time.replace(tzinfo=timezone.utc)
            local_date = observation_time.astimezone(timezone.utc).date().isoformat()
            if observation_time >= cutoff:
                daily_counts[local_date] += 1
        row = {
            "tree": tree,
            "latest": latest,
            "confidence": latest_confidence,
            "uncertain": uncertain,
            "priority": (
                "TINGGI"
                if uncertain or status in {"PERLU_TINDAKAN", "KRITIS"}
                else "NORMAL"
            ),
        }
        tree_rows.append(row)
        if row["priority"] == "TINGGI":
            priority_rows.append(row)
    chart_days = [
        (now.date() - timedelta(days=offset)).isoformat()
        for offset in range(29, -1, -1)
    ]
    chart_values = [daily_counts[day] for day in chart_days]
    max_chart = max(max(chart_values, default=0), 1)
    overall_confidence = (
        sum(confidence_values) / len(confidence_values) if confidence_values else 0.0
    )
    dashboard_metrics = {
        "healthy": status_counts.get("SEHAT", 0),
        "watch": status_counts.get("PERLU_DIAMATI", 0),
        "action": status_counts.get("PERLU_TINDAKAN", 0)
        + status_counts.get("KRITIS", 0),
        "unassessed": status_counts.get("BELUM DIANALISIS", 0),
        "overall_confidence": overall_confidence,
        "priority_count": len(priority_rows),
        "max_chart": max_chart,
    }
    return render_template(
        "dashboard.html",
        trees=trees,
        observation_count=observations,
        tree_rows=tree_rows,
        priority_rows=priority_rows,
        status_counts=dict(status_counts),
        chart_days=chart_days,
        chart_values=chart_values,
        dashboard_metrics=dashboard_metrics,
    )


def _report_row(tree):
    observation = max(
        tree.observations,
        key=lambda item: item.observation_datetime,
        default=None,
    )
    assessment = observation.assessment if observation else None
    return {
        "tree": tree,
        "observation": observation,
        "assessment": assessment,
        "payload": assessment.payload if assessment else None,
    }


def _report_rows():
    trees = db.session.scalars(
        db.select(Tree).where(Tree.active.is_(True)).order_by(Tree.code)
    ).all()
    return [_report_row(tree) for tree in trees]


def _analytics_data(rows):
    assessed = [row for row in rows if row["payload"]]
    confidence_values = []
    risk_levels = Counter()
    status_counts = Counter()
    evidence_counts = Counter()
    for row in assessed:
        payload = row["payload"]
        status = payload.get("agronomic_assessment", {}).get(
            "overall_visual_status", "unknown"
        )
        status_counts[status] += 1
        for item in payload.get("extracted_parameters", []):
            confidence = item.get("confidence")
            if confidence is not None:
                confidence_values.append(float(confidence))
            evidence_counts[item.get("evidence_status", "UNKNOWN")] += 1
        for item in payload.get("agronomic_risks", []):
            risk_levels[item.get("level", "unknown")] += 1
    return {
        "tree_count": len(rows),
        "assessed_count": len(assessed),
        "unassessed_count": len(rows) - len(assessed),
        "average_confidence": (
            sum(confidence_values) / len(confidence_values)
            if confidence_values
            else 0.0
        ),
        "status_counts": dict(status_counts),
        "risk_levels": dict(risk_levels),
        "evidence_counts": dict(evidence_counts),
        "priority_rows": [
            {
                "code": row["tree"].code,
                "status": row["payload"]
                .get("agronomic_assessment", {})
                .get("overall_visual_status", "unknown"),
                "risks": [
                    item
                    for item in row["payload"].get("agronomic_risks", [])
                    if item.get("level", "").lower()
                    in {"tinggi", "high", "sedang", "medium", "potential"}
                ][:5],
            }
            for row in assessed
        ],
    }


@bp.get("/reports")
@login_required
def reports():
    rows = _report_rows()
    return render_template("reports/list.html", rows=rows)


@bp.get("/analytics")
@login_required
def analytics():
    rows = _report_rows()
    return render_template(
        "reports/analytics.html",
        rows=rows,
        analytics=_analytics_data(rows),
    )


@bp.get("/reports/<int:tree_id>")
@login_required
def tree_report(tree_id):
    tree = db.get_or_404(Tree, tree_id)
    row = _report_row(tree)
    if row["assessment"] is None:
        return render_template(
            "reports/detail.html", row=row, payload=None, report_available=False
        )
    return render_template(
        "reports/detail.html",
        row=row,
        payload=row["payload"],
        report_available=True,
    )


@bp.get("/reports/<int:tree_id>.json")
@login_required
def tree_report_json(tree_id):
    tree = db.get_or_404(Tree, tree_id)
    row = _report_row(tree)
    if row["payload"] is None:
        return jsonify(
            {
                "status": "INSUFFICIENT_DATA",
                "tree_id": tree.id,
                "message": "Belum ada assessment v2 untuk pohon ini.",
            }
        ), 404
    return jsonify(row["payload"])


def _export_rows(tree_id=None):
    if tree_id is not None:
        tree = db.get_or_404(Tree, tree_id)
        rows = [_report_row(tree)]
    else:
        rows = _report_rows()
    return [row for row in rows if row["payload"]]


@bp.get("/reports/export.xlsx")
@login_required
def reports_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    rows = _export_rows()
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Ringkasan"
    summary.append(["Kode pohon", "Varietas", "Tanggal observasi", "Status visual", "Versi"])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23613B")
    for row in rows:
        payload = row["payload"]
        summary.append(
            [
                row["tree"].code,
                row["tree"].variety,
                row["observation"].observation_datetime.isoformat()
                if row["observation"]
                else "",
                payload.get("agronomic_assessment", {}).get(
                    "overall_visual_status", "unknown"
                ),
                payload.get("analysis_version", ""),
            ]
        )
    details = workbook.create_sheet("Parameter")
    details.append(["Kode pohon", "Kelompok", "Parameter", "Nilai", "Bukti", "Confidence"])
    for cell in details[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23613B")
    for row in rows:
        for group in (
            "extracted_parameters",
            "architecture_parameters",
            "soil_parameters",
            "weed_parameters",
            "microclimate_parameters",
            "agronomic_risks",
            "plant_status_indices",
        ):
            for item in row["payload"].get(group, []):
                details.append(
                    [
                        row["tree"].code,
                        group,
                        item.get("parameter", item.get("risk", item.get("index", ""))),
                        item.get("value", item.get("status", item.get("level", ""))),
                        item.get("evidence_status", ""),
                        item.get("confidence"),
                    ]
                )
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                max(len(str(cell.value or "")) for cell in column) + 2, 42
            )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="laporan-assessment-rambutan.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.get("/reports/export.pdf")
@login_required
def reports_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    rows = _export_rows()
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Laporan Assessment Agronomi Rambutan v2", styles["Title"]),
        Paragraph(
            "Laporan konservatif berbasis observasi visual; bukan diagnosis dan bukan pengganti pengukuran lapangan.",
            styles["Normal"],
        ),
        Spacer(1, 8 * mm),
    ]
    data = [["Pohon", "Observasi", "Status visual", "Confidence", "Risiko utama"]]
    for row in rows:
        payload = row["payload"]
        risks = payload.get("agronomic_risks", [])
        risk_text = ", ".join(
            str(item.get("risk", "")) for item in risks[:3]
        ) or "Tidak ada"
        confidences = [
            item.get("confidence")
            for item in payload.get("extracted_parameters", [])
            if item.get("confidence") is not None
        ]
        data.append(
            [
                row["tree"].code,
                row["observation"].observation_datetime.strftime("%Y-%m-%d")
                if row["observation"]
                else "-",
                payload.get("agronomic_assessment", {}).get(
                    "overall_visual_status", "unknown"
                ),
                f"{sum(confidences) / len(confidences):.0%}" if confidences else "-",
                Paragraph(risk_text, styles["BodyText"]),
            ]
        )
    table = Table(data, repeatRows=1, colWidths=[28 * mm, 30 * mm, 42 * mm, 25 * mm, 130 * mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#23613B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DCE7DF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8F5")]),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(table)
    document.build(story)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="laporan-assessment-rambutan.pdf",
        mimetype="application/pdf",
    )


def train_models():
    predictions = db.session.scalars(
        db.select(FieldPrediction).where(
            FieldPrediction.field_key == "vision_visual_features"
        )
    ).all()
    corrections = db.session.scalars(db.select(UserCorrection)).all()
    samples = build_training_samples(predictions, corrections)
    result = train_visual_model(samples, current_app.config["MODEL_FOLDER"])
    if result.status == "CANDIDATE_READY":
        metadata = db.session.scalar(
            db.select(ModelMetadata).where(ModelMetadata.model_id == result.model_id)
        )
        if metadata is None:
            metadata = ModelMetadata(
                model_id=result.model_id,
                name="Rambutan visual color classifier",
                task="VISUAL_FEATURE_CLASSIFICATION",
            )
            db.session.add(metadata)
        metadata.version = result.version
        metadata.status = "AVAILABLE"
        metadata.runtime = "scikit-learn-local"
        metadata.checksum = result.checksum
        metadata.metrics_json = json.dumps(
            {
                "accuracy": result.accuracy,
                "samples": result.samples,
                "validation_samples": result.validation_samples,
                "artifact_path": result.artifact_path,
            }
        )
        db.session.commit()
    return jsonify(
        {
            "status": result.status,
            "reason": result.reason,
            "model_id": result.model_id,
            "version": result.version,
            "samples": result.samples,
            "validation_samples": result.validation_samples,
            "accuracy": result.accuracy,
            "artifact_path": result.artifact_path,
            "checksum": result.checksum,
        }
    )


@bp.post("/api/predictions/<int:prediction_id>/correction")
@login_required
def correct_prediction(prediction_id):
    payload = request.get_json(silent=True) or {}
    final_value = str(payload.get("user_final_value", "")).strip()
    if not final_value:
        return (
            jsonify({"status": "INVALID", "message": "Label koreksi wajib diisi."}),
            400,
        )
    prediction = db.get_or_404(FieldPrediction, prediction_id)
    correction = UserCorrection(
        prediction_id=prediction.id,
        ai_original_value=prediction.prediction,
        user_final_value=final_value,
        was_corrected=final_value != prediction.prediction,
        correction_source="USER_CORRECTED",
    )
    db.session.add(correction)
    db.session.commit()
    training_status = train_models().get_json()
    return jsonify(
        {
            "status": "SAVED",
            "message": "Koreksi disimpan sebagai label training; model akan dilatih "
            "setelah dataset memenuhi syarat.",
            "prediction_id": prediction.id,
            "training": training_status,
        }
    )


@bp.get("/backup.json")
@login_required
def backup():
    trees = db.session.scalars(db.select(Tree).order_by(Tree.code)).all()
    export_tables = [
        "agronomic_assessment",
        "growth_measurement",
        "watering_event",
        "fertilizer_event",
        "pruning_event",
        "flowering_observation",
        "fruit_observation",
        "pest_observation",
        "disease_observation",
        "weed_observation",
        "weather_observation",
        "harvest",
        "sensor_device",
        "sensor_reading",
        "model_metadata",
        "system_log",
    ]

    def serialize(value):
        return value.isoformat() if isinstance(value, datetime) else value

    relational_data = {}
    for table_name in export_tables:
        table = db.metadata.tables[table_name]
        rows = db.session.execute(table.select()).mappings()
        relational_data[table_name] = [
            {key: serialize(value) for key, value in row.items()} for row in rows
        ]
    data = {
        "garden": trees[0].garden.name if trees else None,
        "trees": [
            {
                "code": tree.code,
                "species": tree.species,
                "variety": tree.variety,
                "status": tree.status,
                "observations": [
                    {
                        "datetime": observation.observation_datetime.isoformat(),
                        "mode": observation.inspection_mode,
                        "notes": observation.notes,
                        "predictions": [
                            {
                                "field": p.field_key,
                                "value": p.prediction,
                                "confidence": p.confidence,
                            }
                            for p in observation.predictions
                        ],
                    }
                    for observation in tree.observations
                ],
            }
            for tree in trees
        ],
        "relational_data": relational_data,
    }
    payload = build_backup(data)
    response = jsonify(payload)
    response.headers["Content-Disposition"] = (
        "attachment; filename=ai-analis-rambutan-backup.json"
    )
    return response


def _save_photo(file_storage, tree_code):
    original = secure_filename(file_storage.filename or "")
    extension = original.rsplit(".", 1)[-1].lower() if "." in original else ""
    if extension not in current_app.config["ALLOWED_IMAGE_EXTENSIONS"]:
        return None
    folder = Path(current_app.config["UPLOAD_FOLDER"]) / tree_code
    folder.mkdir(parents=True, exist_ok=True)
    filename = f"{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S%f')}.{extension}"
    target = folder / filename
    file_storage.save(target)
    checksum = hashlib.sha256(target.read_bytes()).hexdigest()
    root = Path(current_app.root_path).parent
    try:
        stored_path = target.relative_to(root)
    except ValueError:
        stored_path = target
    return str(stored_path), checksum


def create_inspection(tree, form):
    observation = ObservationSession(
        tree=tree,
        inspection_mode="PHOTO_FIRST",
        observer=current_user.username,
        notes=form.notes.data.strip() if form.notes.data else None,
    )
    db.session.add(observation)
    db.session.flush()
    soil = SoilObservation(
        observation=observation,
        surface_dark=form.surface_dark.data,
        standing_water=form.standing_water.data,
        leaf_wilt=form.leaf_wilt.data,
        soil_moisture_visual=form.soil_moisture_visual.data or None,
        soil_surface_condition=form.soil_surface_condition.data or None,
        soil_compaction=form.soil_compaction.data or None,
        soil_drainage=form.soil_drainage.data or None,
        standing_water_depth_cm=form.standing_water_depth_cm.data,
        visible_cracks=form.visible_cracks.data,
        mulch_present=form.mulch_present.data,
    )
    growth = GrowthMeasurement(
        tree_id=tree.id,
        height_cm=form.height_cm.data,
        stem_diameter_cm=form.stem_diameter_cm.data,
        canopy_width_cm=form.canopy_width_cm.data,
        canopy_height_cm=form.canopy_height_cm.data,
        canopy_ns_cm=form.canopy_ns_cm.data,
        canopy_ew_cm=form.canopy_ew_cm.data,
        primary_branch_count=form.primary_branch_count.data,
        measurement_method="USER_CONFIRMED",
    )
    flowering = FloweringObservation(
        tree_id=tree.id,
        present=form.flower_present.data,
        stage=form.flower_stage.data or None,
    )
    fruit = FruitObservation(
        tree_id=tree.id,
        present=form.fruit_present.data,
        stage=form.fruit_stage.data or None,
        count_estimate=form.fruit_count_estimate.data,
        damage_percent=form.fruit_damage_percent.data,
        fruit_drop_count=form.fruit_drop_count.data,
        fruit_drop_level=form.fruit_drop_level.data or None,
        fruit_color=form.fruit_color.data or None,
        ripeness=form.ripeness.data or None,
    )
    pest = PestObservation(
        tree_id=tree.id,
        present=form.pest_present.data,
        category=form.pest_type.data or None,
        affected_part=form.pest_affected_part.data or None,
        spread=form.pest_spread.data or None,
    )
    disease = DiseaseObservation(
        tree_id=tree.id,
        present=form.disease_present.data,
        category=form.disease_type.data or None,
        affected_part=form.disease_affected_part.data or None,
        spread=form.disease_spread.data or None,
    )
    weed = WeedObservation(
        tree_id=tree.id,
        level=form.weed_level.data or None,
        coverage_percent=form.weed_coverage_percent.data,
        density=form.weed_density.data or None,
        height_cm=form.weed_height_cm.data,
        removed=form.weed_removed.data,
        removal_method=form.weed_removal_method.data or None,
    )
    inputs = InspectionInputs(
        surface_dark=soil.surface_dark,
        standing_water=soil.standing_water,
        leaf_wilt=soil.leaf_wilt,
        photo_count=len(
            [photo for photo in form.photos.data or [] if photo and photo.filename]
        ),
    )
    analysis = analyze_inspection(
        inputs,
        observation_datetime=observation.observation_datetime,
    )
    prediction = analysis.result.predictions[0]
    soil.soil_moisture_visual = prediction.value
    soil.soil_visual_condition = "INFERRED"
    observation.general_condition = analysis.condition
    field_prediction = FieldPrediction(
        observation=observation,
        field_key="root_zone_moisture_class",
        prediction=prediction.value,
        confidence=prediction.confidence,
        source=prediction.source,
        explanation=" ".join(item.text for item in prediction.evidence),
    )
    advanced_prediction = FieldPrediction(
        observation=observation,
        field_key="advanced_explainable_analysis",
        prediction=json.dumps(analysis.advanced, ensure_ascii=False, sort_keys=True),
        confidence=float(analysis.advanced.get("confidence", 0.0)),
        source="HYBRID_RULE_ML_XAI",
        explanation=analysis.advanced.get("recommendation", ""),
    )
    tree.status = analysis.condition
    db.session.add_all(
        [
            soil,
            field_prediction,
            advanced_prediction,
            growth,
            flowering,
            fruit,
            pest,
            disease,
            weed,
        ]
    )
    assessment_payload = None
    if form.ai_analysis.data:
        try:
            auto_analysis = json.loads(form.ai_analysis.data)
        except json.JSONDecodeError:
            auto_analysis = None
        if isinstance(auto_analysis, dict):
            assessment_payload = build_assessment(
                auto_analysis.get("auto_fields", {}),
                manual={
                    "tree_id": tree.id,
                    "height_cm": form.height_cm.data,
                    "stem_diameter_cm": form.stem_diameter_cm.data,
                    "canopy_width_cm": form.canopy_width_cm.data,
                    "flower_present": form.flower_present.data,
                    "fruit_present": form.fruit_present.data,
                    "standing_water": form.standing_water.data,
                    "leaf_wilt": form.leaf_wilt.data,
                    "pest_present": form.pest_present.data,
                    "disease_present": form.disease_present.data,
                    "weed_level": form.weed_level.data,
                    "soil_surface_condition": form.soil_surface_condition.data,
                    "soil_compaction": form.soil_compaction.data,
                    "mulch_present": form.mulch_present.data,
                },
                photo_count=len(form.photos.data or []),
            )
            db.session.add(
                FieldPrediction(
                    observation=observation,
                    field_key="vision_auto_extraction",
                    prediction=json.dumps(
                        auto_analysis, ensure_ascii=False, sort_keys=True
                    ),
                    confidence=0.0,
                    source="COMPUTER_VISION",
                    explanation=(
                        "Ringkasan ekstraksi otomatis sebelum penyimpanan; "
                        "nilai agronomi yang tidak didukung model tetap perlu konfirmasi."
                    ),
                )
            )
            db.session.add(
                AgronomicAssessment(
                    observation=observation,
                    analysis_type=assessment_payload["analysis_type"],
                    analysis_version=assessment_payload["analysis_version"],
                    overall_status=assessment_payload["agronomic_assessment"]["overall_visual_status"],
                    confirmation_required=assessment_payload["quality_control"]["user_confirmation_required"],
                    payload_json=json.dumps(assessment_payload, ensure_ascii=False, sort_keys=True),
                )
            )
    if assessment_payload is None:
        assessment_payload = build_assessment(
            manual={
                "tree_id": tree.id,
                "height_cm": form.height_cm.data,
                "stem_diameter_cm": form.stem_diameter_cm.data,
                "canopy_width_cm": form.canopy_width_cm.data,
                "flower_present": form.flower_present.data,
                "fruit_present": form.fruit_present.data,
                "standing_water": form.standing_water.data,
                "leaf_wilt": form.leaf_wilt.data,
                "pest_present": form.pest_present.data,
                "disease_present": form.disease_present.data,
                "weed_level": form.weed_level.data,
                "soil_surface_condition": form.soil_surface_condition.data,
                "soil_compaction": form.soil_compaction.data,
                "mulch_present": form.mulch_present.data,
            }
        )
        db.session.add(
            AgronomicAssessment(
                observation=observation,
                analysis_type=assessment_payload["analysis_type"],
                analysis_version=assessment_payload["analysis_version"],
                overall_status=assessment_payload["agronomic_assessment"]["overall_visual_status"],
                confirmation_required=True,
                payload_json=json.dumps(assessment_payload, ensure_ascii=False, sort_keys=True),
            )
        )
    db.session.flush()
    for file_storage in form.photos.data or []:
        if not file_storage or not file_storage.filename:
            continue
        saved = _save_photo(file_storage, tree.code)
        if saved:
            path, checksum = saved
            photo_path = Path(path)
            if not photo_path.is_absolute():
                photo_path = Path(current_app.root_path).parent / photo_path
            quality = assess_image(photo_path)
            db.session.add(
                Photo(
                    observation=observation,
                    capture_type="INSPECTION",
                    file_path=path,
                    quality_status=quality.status,
                    quality_reasons=" ".join(quality.reasons),
                    checksum=checksum,
                )
            )
            model_path = (
                Path(
                    current_app.config.get(
                        "MODEL_FOLDER",
                        Path(current_app.root_path).parent / "instance" / "models",
                    )
                )
                / "yolo11n.pt"
            )
            if quality.status == "GOOD" and model_path.exists():
                vision_results = analyze_available_models(
                    photo_path,
                    current_app.config.get(
                        "MODEL_FOLDER",
                        Path(current_app.root_path).parent / "instance" / "models",
                    ),
                )
                vision_result = vision_results["object_detection"]
                if vision_result["status"] == "READY":
                    detections = vision_result.get("detections", [])
                    max_confidence = max(
                        (item["confidence"] for item in detections), default=0.0
                    )
                for model_key, model_result in vision_results.items():
                    if model_key == "object_detection":
                        continue
                    if model_result["status"] in {"READY", "INSUFFICIENT_DATA"}:
                        db.session.add(
                            FieldPrediction(
                                observation=observation,
                                field_key=f"vision_{model_key}",
                                prediction=json.dumps(
                                    _analysis_summary({model_key: model_result})[
                                        model_key
                                    ],
                                    ensure_ascii=False,
                                    sort_keys=True,
                                ),
                                confidence=float(model_result.get("confidence", 0.0)),
                                source="COMPUTER_VISION",
                                explanation=(
                                    model_result.get("message")
                                    or "Hasil model visual disimpan sebagai bukti teknis."
                                ),
                            )
                        )
                    db.session.add(
                        FieldPrediction(
                            observation=observation,
                            field_key="vision_object_detection",
                            prediction=json.dumps(
                                detections, ensure_ascii=False, sort_keys=True
                            ),
                            confidence=max_confidence,
                            source="COMPUTER_VISION",
                            explanation=(
                                "Deteksi objek YOLO lokal. " "Bukan diagnosis agronomi."
                            ),
                        )
                    )
    db.session.commit()
    inference = type(
        "SoilInferenceView",
        (),
        {
            "moisture": prediction.value,
            "confidence": prediction.confidence,
            "explanations": [item.text for item in prediction.evidence],
        },
    )()
    return observation, inference, analysis.condition, analysis.condition_reasons
